from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
CAP = ROOT / "Main_GUI_캡처"
BEFORE = CAP / "Main_GUI_프로그램_시작전.png"
AFTER = CAP / "Main_GUI_프로그램_시작후.png"
TRAP = CAP / "Main_GUI_프로그램_시작후_Trap수신.png"
DETAIL = CAP / "Main_GUI_프로그램_전체모듈상세정보.png"
ORDER = CAP / "Main_GUI_프로그램_모듈설치순서설정.png"
ORDER_LIST = CAP / "Main_GUI_프로그램_모듈설치순서설정_목록.png"
ORDER_SET = CAP / "Main_GUI_프로그램_모듈설치순서설정_설정.png"
ORDER_AFTER = CAP / "Main_GUI_프로그램_모듈설치순서설정후_메인화면.png"
RECORD = CAP / "Main_GUI_운전데이타기록설정.png"
COMPANY_LOGO = ROOT / "pantech-original.png"
CUSTOMER_LOGO = ROOT / "skt_logo-original.png"
OUTPUT = ROOT / "TBC1000B_MainGUI_고객사용자매뉴얼.pptx"

NAVY = RGBColor(12, 52, 111)
BLUE = RGBColor(38, 105, 255)
PALE_BLUE = RGBColor(238, 244, 255)
GREEN = RGBColor(3, 152, 85)
PALE_GREEN = RGBColor(236, 253, 243)
RED = RGBColor(217, 45, 32)
PALE_RED = RGBColor(254, 243, 242)
ORANGE = RGBColor(247, 144, 9)
YELLOW = RGBColor(247, 201, 72)
GRAY = RGBColor(71, 84, 103)
MID_GRAY = RGBColor(208, 213, 221)
LIGHT_GRAY = RGBColor(242, 244, 247)
WHITE = RGBColor(255, 255, 255)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def text(slide, value, x, y, w, h, size=16, color=GRAY, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=None, rounded=True, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.fill.transparency = transparency
    s.line.color.rgb = line or fill
    return s


def header(slide, title, section, page):
    rect(slide, 0, 0, 13.333, 0.12, BLUE, rounded=False)
    text(slide, title, 0.55, 0.24, 10.6, 0.52, 24, NAVY, True)
    text(slide, section, 10.7, 0.29, 1.95, 0.3, 9, BLUE, True, PP_ALIGN.RIGHT)
    rect(slide, 0.55, 0.83, 12.2, 0.012, MID_GRAY, rounded=False)
    text(slide, "TBC1000B Main GUI  |  고객 사용자 매뉴얼", 0.55, 7.06, 7.2, 0.18, 8,
         RGBColor(152, 162, 179))
    text(slide, str(page), 12.55, 7.06, 0.3, 0.18, 8, RGBColor(152, 162, 179),
         align=PP_ALIGN.RIGHT)


def bullets(slide, values, x, y, w, h, size=14, color=GRAY, spacing=7):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, value in enumerate(values):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + value
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
    return shape


def step(slide, n, title, body, x, y, w=5.25, color=BLUE):
    rect(slide, x, y, w, 1.03, WHITE, MID_GRAY)
    b = rect(slide, x + 0.18, y + 0.22, 0.52, 0.52, color, color)
    tf = b.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"STEP {n}" if isinstance(n, int) else str(n)
    r.font.name = FONT
    r.font.size = Pt(9 if isinstance(n, int) else 15)
    r.font.bold = True
    r.font.color.rgb = WHITE
    text(slide, title, x + 0.86, y + 0.1, w - 1.04, 0.32, 15, NAVY, True)
    text(slide, body, x + 0.86, y + 0.44, w - 1.04, 0.43, 10, GRAY)


def crop_stream(path, crop=None):
    with Image.open(path) as im:
        image = im.crop(crop) if crop else im.copy()
        size = image.size
        stream = BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
    return stream, size


def picture(slide, path, x, y, w, h, crop=None, border=True):
    stream, (iw, ih) = crop_stream(path, crop)
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    px, py = x + (w - dw) / 2, y + (h - dh) / 2
    slide.shapes.add_picture(stream, Inches(px), Inches(py), Inches(dw), Inches(dh))
    if border:
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px), Inches(py), Inches(dw), Inches(dh))
        s.fill.background()
        s.line.color.rgb = MID_GRAY
        s.line.width = Pt(1)
    return px, py, dw, dh, iw, ih


def annotated_picture(slide, path, x, y, w, h, marks, crop=None):
    px, py, dw, dh, iw, ih = picture(slide, path, x, y, w, h, crop)
    for n, (mx, my, mw, mh), color in marks:
        rx, ry = px + dw * mx / iw, py + dh * my / ih
        rw, rh = dw * mw / iw, dh * mh / ih
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(rx), Inches(ry), Inches(rw), Inches(rh))
        s.fill.background()
        s.line.color.rgb = color
        s.line.width = Pt(2.5)
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(rx - 0.1), Inches(ry - 0.1), Inches(0.32), Inches(0.32))
        c.fill.solid(); c.fill.fore_color.rgb = color; c.line.color.rgb = WHITE
        tf = c.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(n); r.font.name = FONT; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = WHITE


def note(slide, title, body, x, y, w, h, color=BLUE, pale=PALE_BLUE):
    rect(slide, x, y, w, h, pale, color)
    text(slide, title, x + 0.23, y + 0.1, w - 0.46, 0.28, 12, color, True)
    text(slide, body, x + 0.23, y + 0.38, w - 0.46, h - 0.46, 10, GRAY,
         valign=MSO_ANCHOR.TOP)


# 1 Cover
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, RGBColor(247, 249, 252), rounded=False)
rect(s, 0, 0, 4.1, 7.5, NAVY, rounded=False)
text(s, "TBC1000B", 0.58, 1.1, 3.0, 0.55, 25, WHITE, True)
text(s, "Battery Monitoring System", 0.58, 1.72, 3.0, 0.75, 19, RGBColor(180, 205, 255), True)
text(s, "Main GUI", 0.58, 2.55, 3.0, 0.45, 18, WHITE, True)
text(s, "고객 사용자 매뉴얼", 4.82, 2.02, 7.7, 0.75, 34, NAVY, True)
text(s, "시스템 접속 · 상태 모니터링 · 알람 · 데이터 기록 · EPO", 4.85, 2.92, 7.5, 0.42, 15, GRAY)
rect(s, 4.85, 3.58, 1.05, 0.07, BLUE, rounded=False)
text(s, "Base SNMPv2  v3.2.1", 4.85, 3.88, 3.2, 0.34, 13, BLUE, True)
text(s, "2026. 08", 4.85, 4.3, 2.0, 0.28, 11, RGBColor(102, 112, 133))
rect(s, 4.85, 5.14, 7.65, 1.12, WHITE, MID_GRAY)
text(s, "CUSTOMER", 5.08, 5.22, 1.4, 0.2, 8, RGBColor(152, 162, 179), True)
picture(s, CUSTOMER_LOGO, 5.08, 5.43, 2.05, 0.64, border=False)
rect(s, 8.35, 5.34, 0.012, 0.72, MID_GRAY, rounded=False)
text(s, "SYSTEM ENGINEERING", 8.67, 5.22, 2.0, 0.2, 8, RGBColor(152, 162, 179), True)
picture(s, COMPANY_LOGO, 9.42, 5.4, 2.72, 0.7, border=False)
text(s, "EPO 및 설정 변경 기능은 승인된 운영자만 사용하십시오.", 4.85, 6.48, 7.5, 0.34, 11, RED, True)

# 2 Quick start
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "사용 절차 한눈에 보기", "QUICK START", 2)
items = [
    (1, "접속 설정 확인", "IP, Port, GET/SET/TRAP Community와 TRAP Port 확인"),
    (2, "설치 정보 입력", "설치 장소·축전지명·설비 정보 입력 후 저장"),
    (3, "접속 시작", "접속상태 녹색 표시와 최종 업데이트 시각 확인"),
    (4, "상태 모니터링", "시스템 요약, 모듈 상태, 고장 정보 확인"),
    (5, "알람·상세 확인", "TRAP 로그, 알람 목록, 전체모듈정보 확인"),
    (6, "기록·제어", "운전 데이터 기록 및 승인된 설정/EPO 수행"),
]
for i, (n, t, b) in enumerate(items):
    step(s, n, t, b, 0.72 + (i % 2) * 6.2, 1.2 + (i // 2) * 1.52, 5.72,
         GREEN if n == 4 else BLUE)
note(s, "안전 원칙", "차단·복구 및 충전 제한값 변경 전 부하 상태와 현장 작업 승인을 확인합니다.",
     0.72, 5.92, 11.92, 0.72, RED, PALE_RED)

# 3 Main screen map
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "1. 메인 화면 구성", "OVERVIEW", 3)
annotated_picture(s, AFTER, 0.55, 1.03, 9.15, 5.85, [
    ("A", (6, 29, 1845, 83), BLUE),
    ("B", (6, 120, 1845, 81), GREEN),
    ("C", (16, 215, 590, 240), ORANGE),
    ("D", (612, 215, 1225, 235), BLUE),
    ("E", (6, 467, 1560, 370), GREEN),
    ("F", (1571, 467, 278, 370), ORANGE),
    ("G", (6, 850, 1842, 214), RED),
])
text(s, "화면 영역 안내", 9.98, 1.08, 2.5, 0.38, 18, NAVY, True)
area_legend = (
    ("A", "시스템 접속 설정", BLUE),
    ("B", "설치 및 설비 정보", GREEN),
    ("C", "시스템 요약 정보", ORANGE),
    ("D", "SNMP Trap 로그", BLUE),
    ("E", "모듈 상태", GREEN),
    ("F", "모듈 설치 순서", ORANGE),
    ("G", "고장 정보", RED),
)
for i, (letter, label, color) in enumerate(area_legend):
    y = 1.62 + i * 0.69
    badge = rect(s, 9.98, y, 0.46, 0.42, color, color)
    tf = badge.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = letter
    r.font.name = FONT
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    text(s, label, 10.58, y - 0.01, 2.05, 0.43, 11, GRAY, True)

# 4 Connection
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "2. 시스템 접속 및 설치 정보 저장", "CONNECT", 4)
annotated_picture(s, BEFORE, 0.55, 1.08, 7.15, 2.15, [
    (1, (5, 25, 1847, 88), BLUE),
    (2, (5, 119, 1847, 81), GREEN),
], crop=(0, 0, 1858, 215))
step(s, 1, "접속 값 확인", "IP·Port와 GET/SET/TRAP Community, TRAP Port를 확인합니다.", 7.95, 1.12, 4.72)
step(s, 2, "설치 정보 입력", "설치 장소, 축전지명 및 설비 정보를 입력하고 [저장]을 클릭합니다.", 7.95, 2.28, 4.72, GREEN)
step(s, 3, "접속 시작", "[접속시작] 클릭 후 접속상태가 녹색인지 확인합니다.", 0.7, 3.62, 5.72)
step(s, 4, "최종 업데이트", "수신 시각이 현재 시각에 맞게 계속 갱신되는지 확인합니다.", 6.68, 3.62, 5.72, GREEN)
note(s, "접속 실패 시", "대상 장비 전원, 네트워크 연결, IP/Port, Community 값을 점검한 뒤 다시 접속합니다.",
     0.7, 5.07, 11.7, 0.78, ORANGE, RGBColor(255, 250, 235))
note(s, "보안 주의", "Community 문자열은 접속 권한 정보입니다. 화면 캡처나 외부 전달 자료에서 노출되지 않도록 관리하십시오.",
     0.7, 6.02, 11.7, 0.66, RED, PALE_RED)

# 5 System summary
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "3. 시스템 요약 정보 확인", "MONITOR", 5)
annotated_picture(s, AFTER, 0.55, 1.08, 6.55, 5.65, [
    (1, (0, 0, 590, 225), BLUE),
    (2, (465, 125, 125, 48), RED),
], crop=(12, 215, 606, 455))
text(s, "주요 확인 항목", 7.48, 1.2, 4.7, 0.42, 21, NAVY, True)
bullets(s, [
    "Rack 전압, 전류와 SOC 충전율",
    "Max/Min/Avg 전압 및 온도",
    "방전 횟수와 충전 제한 설정값",
    "과전압·저온·과전류 차단 상태",
    "설비번호, 운영 관리자, 제조사 등 기본정보",
], 7.42, 1.8, 5.05, 2.75, 14)
note(s, "색상 확인", "녹색은 정상 상태입니다. 비정상 색상이나 '-' 값이 보이면 접속 상태와 해당 항목을 점검합니다.",
     7.38, 4.72, 5.1, 0.95, GREEN, PALE_GREEN)
note(s, "EPO", "[차단]/[복구]는 전체 모듈 제어 기능입니다. 작업 승인 후에만 사용합니다.",
     7.38, 5.85, 5.1, 0.82, RED, PALE_RED)

# 6 Module state
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "4. 모듈 상태 및 고장 정보 확인", "MONITOR", 6)
annotated_picture(s, AFTER, 0.55, 1.05, 8.1, 5.92, [
    (1, (0, 0, 1560, 380), BLUE),
    (2, (0, 382, 1840, 215), RED),
], crop=(5, 465, 1848, 1068))
text(s, "1  모듈 상태 표", 8.98, 1.18, 3.5, 0.35, 19, BLUE, True)
bullets(s, [
    "모듈 전압",
    "셀 전압 Max/Min",
    "셀 온도 Max/Min",
    "경보 및 통신상태",
    "상세정보와 개별 EPO",
], 8.92, 1.66, 3.65, 2.15, 13)
text(s, "상태 표의 경보 색상 기준", 8.98, 3.92, 3.4, 0.35, 16, NAVY, True)
for i, (name, color) in enumerate((("Critical", RED), ("Major", ORANGE), ("Minor", YELLOW), ("Warning", RGBColor(116, 192, 252)))):
    rect(s, 8.95 + (i % 2) * 1.78, 4.46 + (i // 2) * 0.52, 1.55, 0.38, color, color)
    text(s, name, 8.95 + (i % 2) * 1.78, 4.47 + (i // 2) * 0.52, 1.55, 0.34, 10,
         WHITE if i < 2 else NAVY, True, PP_ALIGN.CENTER)
note(s, "2  고장 정보", "하단 표에서 고장 모듈·셀 번호와 전압·온도를 확인합니다. 원인 조치 전 임의 삭제하지 마십시오.",
     8.88, 5.65, 3.8, 1.0, RED, PALE_RED)

# 7 Module detail
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "5. 전체 모듈 상세정보", "DETAIL", 7)
picture(s, DETAIL, 0.55, 1.05, 12.2, 4.85)
note(s, "열기", "메인 화면의 [전체모듈정보]를 클릭합니다.", 0.7, 6.05, 2.75, 0.62, BLUE, PALE_BLUE)
note(s, "식별 정보", "모듈 번호, SW버전, Equip ID, 모델, 바코드를 확인합니다.", 3.62, 6.05, 3.4, 0.62, GREEN, PALE_GREEN)
note(s, "운전 정보", "전압·전류·상태·SOC/SOH와 각 셀 전압·온도를 좌우 스크롤로 확인합니다.",
     7.18, 6.05, 5.35, 0.62, ORANGE, RGBColor(255, 250, 235))

# 8 Module order selection
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "6. 모듈 설치 순서 설정 — 모듈 선택", "CONFIGURATION", 8)
annotated_picture(s, ORDER_LIST, 0.72, 1.05, 5.45, 5.95, [
    (1, (13, 42, 95, 330), BLUE),
    (2, (207, 86, 424, 226), GREEN),
    (3, (329, 706, 310, 24), ORANGE),
])
text(s, "선택 절차", 6.6, 1.2, 5.4, 0.42, 21, NAVY, True)
step(s, 1, "실제 위치 확인", "좌측 Rack 그림의 01~10 위치를 기준으로 확인합니다.", 6.52, 1.84, 5.75)
step(s, 2, "모듈 선택", "드롭다운 목록을 펼치고 모듈 번호와 바코드를 확인하여 선택합니다.", 6.52, 3.02, 5.75, GREEN)
step(s, 3, "설정 저장", "필요한 위치를 모두 지정한 뒤 [저장]을 클릭합니다.", 6.52, 4.20, 5.75, ORANGE)
note(s, "선택 기준", "바코드로 실물 모듈을 대조하십시오. 같은 모듈은 두 위치에 중복 선택할 수 없습니다.",
     6.52, 5.58, 5.75, 0.88, ORANGE, RGBColor(255, 250, 235))

# 9 Module order save
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "6. 모듈 설치 순서 설정 — 확인 및 저장", "CONFIGURATION", 9)
annotated_picture(s, ORDER_SET, 0.72, 1.05, 5.45, 5.95, [
    (1, (207, 88, 424, 276), GREEN),
    (2, (329, 706, 310, 24), BLUE),
    (3, (12, 706, 310, 24), ORANGE),
])
text(s, "저장 전 확인", 6.6, 1.2, 5.4, 0.42, 21, NAVY, True)
step(s, 1, "전체 위치 대조", "10번부터 1번까지 모듈·바코드의 중복과 누락을 확인합니다.", 6.52, 1.84, 5.75, GREEN)
step(s, 2, "저장", "설정이 정확하면 [저장]을 클릭하여 설치 순서를 적용합니다.", 6.52, 3.02, 5.75, BLUE)
step(s, 3, "초기화", "전체 재설정이 필요한 경우에만 [초기화]를 사용합니다.", 6.52, 4.20, 5.75, ORANGE)
note(s, "초기화 주의", "초기화를 실행하면 저장된 설치 순서가 모두 해제되므로 작업 전에 반드시 확인하십시오.",
     6.52, 5.58, 5.75, 0.88, ORANGE, RGBColor(255, 250, 235))

# 10 Module order verification
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "6. 모듈 설치 순서 설정 — 적용 결과 확인", "CONFIGURATION", 10)
annotated_picture(s, ORDER_AFTER, 0.55, 1.05, 8.35, 5.92, [
    (1, (1570, 500, 280, 330), BLUE),
])
text(s, "1  메인 화면 설치 순서", 9.2, 1.2, 3.3, 0.4, 18, BLUE, True)
bullets(s, [
    "Rack 10번부터 1번까지 순서 확인",
    "각 위치의 바코드 확인",
    "괄호 안 모듈 번호 확인",
    "빨간색: 알람 발생",
    "갈색: 바코드 불일치",
], 9.12, 1.78, 3.45, 2.5, 13)
note(s, "정상 적용", "저장한 위치·바코드·모듈 번호가 우측 '모듈 설치 순서'에 동일하게 표시되어야 합니다.",
     9.08, 4.62, 3.58, 1.0, GREEN, PALE_GREEN)
note(s, "불일치 시", "실물 바코드를 다시 확인한 후 [모듈 설치 순서 설정]에서 해당 위치를 수정합니다.",
     9.08, 5.82, 3.58, 0.82, RED, PALE_RED)

# 11 Alarm and trap
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "7. 알람 및 SNMP Trap 확인", "ALARM", 11)
annotated_picture(s, TRAP, 0.55, 1.05, 8.15, 5.92, [
    (1, (620, 245, 1210, 170), BLUE),
    (2, (718, 128, 193, 31), RED),
    (3, (829, 224, 85, 27), GREEN),
])
text(s, "1  TRAP 로그", 9.02, 1.2, 3.4, 0.38, 20, BLUE, True)
bullets(s, [
    "수신 시간과 Trap OID",
    "Alarm 및 Level",
    "Equip ID/Name",
    "최대 1,000개 로그 저장",
], 8.96, 1.73, 3.55, 1.75, 13)
note(s, "2  발생된 알람 보기", "괄호 안 숫자는 현재 발생 알람 수입니다. 버튼을 눌러 목록을 확인합니다.",
     8.92, 3.72, 3.72, 0.9, RED, PALE_RED)
note(s, "3  재전송 요청", "누락이 의심되는 경우에만 [재전송요청]을 사용합니다.",
     8.92, 4.83, 3.72, 0.78, GREEN, PALE_GREEN)
note(s, "로그 삭제", "확인 및 기록 완료 후에만 TRAP 로그 전체 삭제를 실행합니다.",
     8.92, 5.82, 3.72, 0.78, ORANGE, RGBColor(255, 250, 235))

# 12 Alarm options
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "8. 알람 표시 및 수신 옵션", "ALARM", 12)
picture(s, AFTER, 0.65, 1.15, 12.0, 1.05, crop=(1000, 30, 1515, 112))
text(s, "알람 레벨 선택", 0.78, 2.48, 5.2, 0.4, 21, NAVY, True)
bullets(s, [
    "CRIT: Critical 알람",
    "MAJOR: Major 알람",
    "MINOR: Minor 알람",
    "WARN: Warning 알람",
    "체크된 등급을 기준으로 알림을 사용합니다.",
], 0.75, 3.0, 5.45, 2.25, 14)
text(s, "상태 표시", 6.85, 2.48, 5.2, 0.4, 21, NAVY, True)
bullets(s, [
    "접속상태: 녹색이면 정상 접속",
    "Timeout: 통신 실패 누적 상태 확인",
    "스피커 아이콘: 알람 음향 상태",
    "TX/RX: 송수신 동작 표시",
    "CPU/MEM/TRAP/QUEUE: 프로그램 처리 상태",
], 6.82, 3.0, 5.65, 2.25, 14)
note(s, "운영 권고", "CRIT·MAJOR 알람은 항상 활성화하고, 현장 정책에 따라 MINOR·WARN 사용 여부를 결정합니다.",
     0.75, 5.72, 11.72, 0.82, RED, PALE_RED)

# 13 Recording
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "9. 운전 데이터 기록", "DATA RECORD", 13)
annotated_picture(s, RECORD, 0.72, 1.05, 4.25, 5.95, [
    (1, (18, 77, 354, 57), BLUE),
    (2, (18, 180, 354, 178), GREEN),
    (3, (18, 364, 354, 133), ORANGE),
    (4, (18, 537, 354, 29), GREEN),
])
text(s, "기록 설정 절차", 5.35, 1.18, 6.7, 0.42, 21, NAVY, True)
step(s, 1, "주기와 기간", "기록 간격 및 기록 기간을 선택합니다.", 5.28, 1.78, 6.65)
step(s, 2, "Excel 항목 선택", "필요 항목을 체크합니다. 기록시각과 모듈 번호는 항상 저장됩니다.", 5.28, 2.94, 6.65, GREEN)
step(s, 3, "용량 확인", "예상 기록량, 권장 공간, 현재 여유 공간과 저장 위치를 확인합니다.", 5.28, 4.10, 6.65, ORANGE)
step(s, 4, "기록 시작/중지", "[기록시작] 후 상태 문구를 확인하고, 종료 시 [기록중지]를 누릅니다.", 5.28, 5.26, 6.65, GREEN)

# 14 EPO and controls
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "10. EPO 및 충전 제한 제어", "CONTROL", 14)
annotated_picture(s, AFTER, 0.55, 1.05, 7.25, 3.15, [
    (1, (456, 110, 136, 49), RED),
    (2, (387, 178, 205, 36), BLUE),
], crop=(20, 215, 610, 450))
text(s, "제어 전 확인", 8.12, 1.18, 4.35, 0.38, 20, NAVY, True)
bullets(s, [
    "현재 부하와 축전지 운전 상태",
    "대상 시스템과 모듈 번호",
    "현장 작업 승인 및 안전 조치",
    "확인 팝업의 작업 종류",
], 8.06, 1.72, 4.45, 1.85, 13)
note(s, "1  전체 EPO", "시스템 요약의 [차단]/[복구]는 전체 모듈에 명령을 전송합니다.",
     0.72, 4.52, 5.85, 0.82, RED, PALE_RED)
note(s, "개별 EPO", "모듈 상태 표의 [차단]은 선택한 모듈에만 적용됩니다.",
     6.75, 4.52, 5.85, 0.82, ORANGE, RGBColor(255, 250, 235))
note(s, "2  충전 제한", "충전전류제한(C)과 SOC충전제한(%)은 승인된 운전 기준값만 입력합니다.",
     0.72, 5.58, 5.85, 0.82, BLUE, PALE_BLUE)
note(s, "완료 확인", "진행 상태와 성공/실패 모듈을 확인하고 실패 시 모듈 번호를 기록합니다.",
     6.75, 5.58, 5.85, 0.82, GREEN, PALE_GREEN)

# 15 Finish / troubleshooting
s = prs.slides.add_slide(prs.slide_layouts[6]); header(s, "11. 접속 종료 및 문제 해결", "FINISH", 15)
text(s, "정상 종료", 0.75, 1.2, 5.4, 0.4, 21, NAVY, True)
step(s, 1, "기록·제어 종료", "운전 데이터 기록과 진행 중 제어 작업이 없는지 확인합니다.", 0.7, 1.82, 5.72)
step(s, 2, "접속 종료", "[접속종료]를 클릭하고 접속상태가 대기로 바뀌는지 확인합니다.", 0.7, 3.0, 5.72)
step(s, 3, "프로그램 종료", "데이터 저장 완료 후 창을 닫습니다.", 0.7, 4.18, 5.72)
text(s, "문제 발생 시 확인", 6.88, 1.2, 5.4, 0.4, 21, NAVY, True)
issues = [
    ("접속 실패", "전원·네트워크·IP/Port·Community 확인"),
    ("값 미갱신", "접속상태·Timeout·최종 업데이트 시각 확인"),
    ("TRAP 미수신", "TRAP Community/Port와 방화벽 확인"),
    ("모듈 불일치", "설치 순서와 모듈 바코드 재확인"),
    ("EPO 실패", "실패 모듈 번호 및 처리시각 기록"),
]
for i, (title, body) in enumerate(issues):
    y = 1.82 + i * 0.84
    rect(s, 6.82, y, 5.78, 0.66, LIGHT_GRAY, MID_GRAY)
    text(s, title, 7.08, y + 0.09, 1.18, 0.28, 12, NAVY, True)
    text(s, body, 8.35, y + 0.06, 3.95, 0.38, 10, GRAY)
note(s, "문의 정보", "화면 캡처, 발생 시각, 시스템 IP, 알람/실패 모듈 번호를 함께 전달하십시오.",
     0.7, 5.72, 11.9, 0.78, BLUE, PALE_BLUE)


prs.core_properties.title = "TBC1000B Main GUI 고객 사용자 매뉴얼"
prs.core_properties.subject = "접속, 모니터링, 알람, 데이터 기록 및 EPO 사용 방법"
prs.core_properties.author = "PANTECH C&I Engineering"
prs.core_properties.keywords = "TBC1000B, Main GUI, SNMP, TRAP, EPO, 사용자 매뉴얼"
prs.save(OUTPUT)

# Re-open and test the OPC ZIP package so a damaged handoff is never produced.
check = Presentation(str(OUTPUT))
assert len(check.slides) == 15
with ZipFile(OUTPUT) as zf:
    assert zf.testzip() is None
print(OUTPUT)
print(f"slides={len(check.slides)}, bytes={OUTPUT.stat().st_size}")
