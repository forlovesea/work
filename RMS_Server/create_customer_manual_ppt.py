from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
BEFORE = ROOT / "mini_시작전_프로그램.png"
AFTER = ROOT / "mini_시작후_프로그램.png"
OUTPUT = ROOT / "TBC1000B_MiniGUI_고객사용자매뉴얼(260810).pptx"

NAVY = RGBColor(12, 52, 111)
BLUE = RGBColor(38, 105, 255)
PALE_BLUE = RGBColor(238, 244, 255)
GREEN = RGBColor(3, 152, 85)
PALE_GREEN = RGBColor(236, 253, 243)
RED = RGBColor(217, 45, 32)
PALE_RED = RGBColor(254, 243, 242)
ORANGE = RGBColor(247, 144, 9)
GRAY = RGBColor(71, 84, 103)
LIGHT_GRAY = RGBColor(242, 244, 247)
MID_GRAY = RGBColor(208, 213, 221)
WHITE = RGBColor(255, 255, 255)
FONT = "맑은 고딕"


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_text(slide, text, x, y, w, h, size=18, color=GRAY, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=True, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    return shape


def add_header(slide, title, section=None, page=None):
    add_rect(slide, 0, 0, 13.333, 0.13, BLUE, radius=False)
    add_text(slide, title, 0.55, 0.25, 10.8, 0.52, 25, NAVY, True)
    if section:
        add_text(slide, section, 10.7, 0.28, 1.9, 0.35, 10, BLUE, True, PP_ALIGN.RIGHT)
    if page is not None:
        add_text(slide, str(page), 12.55, 7.06, 0.3, 0.2, 9, RGBColor(152, 162, 179), False, PP_ALIGN.RIGHT)
    add_rect(slide, 0.55, 0.84, 12.2, 0.015, MID_GRAY, radius=False)


def add_footer(slide):
    add_text(slide, "TBC1000B Battery Control Mini  |  고객 사용자 매뉴얼", 0.55, 7.05, 7.0, 0.2, 8,
             RGBColor(152, 162, 179))


def add_bullet_list(slide, items, x, y, w, h, size=17, color=GRAY, spacing=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.level = 0
        p.text = "•  " + p.text
    return box


def add_image_contain(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    px, py = x + (w - dw) / 2, y + (h - dh) / 2
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(dw), Inches(dh))
    return px, py, dw, dh, iw, ih


def add_image_with_callouts(slide, path, x, y, w, h, callouts):
    px, py, dw, dh, iw, ih = add_image_contain(slide, path, x, y, w, h)
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px), Inches(py), Inches(dw), Inches(dh))
    border.fill.background()
    border.line.color.rgb = MID_GRAY
    border.line.width = Pt(1)
    for number, (cx, cy, cw, ch), color in callouts:
        rx = px + dw * cx / iw
        ry = py + dh * cy / ih
        rw = dw * cw / iw
        rh = dh * ch / ih
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(rx), Inches(ry), Inches(rw), Inches(rh))
        rect.fill.background()
        rect.line.color.rgb = color
        rect.line.width = Pt(2.5)
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(rx - 0.12), Inches(ry - 0.12), Inches(0.34), Inches(0.34))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.color.rgb = WHITE
        badge.line.width = Pt(1)
        tf = badge.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(number)
        r.font.name = FONT
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = WHITE


def add_step_card(slide, number, title, body, x, y, w=5.3, color=BLUE):
    add_rect(slide, x, y, w, 1.05, WHITE, MID_GRAY)
    badge = add_rect(slide, x + 0.18, y + 0.22, 0.54, 0.54, color, color)
    tf = badge.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = WHITE
    add_text(slide, title, x + 0.88, y + 0.12, w - 1.05, 0.3, 16, NAVY, True)
    add_text(slide, body, x + 0.88, y + 0.46, w - 1.05, 0.42, 11, GRAY)


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.333, 7.5, RGBColor(247, 249, 252), radius=False)
add_rect(slide, 0, 0, 4.15, 7.5, NAVY, radius=False)
add_text(slide, "TBC1000B", 0.58, 1.2, 3.0, 0.55, 25, WHITE, True)
add_text(slide, "Battery Control Mini", 0.58, 1.82, 3.0, 0.55, 20, RGBColor(180, 205, 255), True)
add_text(slide, "고객 사용자 매뉴얼", 4.85, 2.22, 7.6, 0.75, 34, NAVY, True)
add_text(slide, "정류기용 축전지 시스템의 접속, 상태 확인 및 제어", 4.88, 3.05, 7.1, 0.42, 16, GRAY)
add_rect(slide, 4.88, 3.72, 1.05, 0.07, BLUE, radius=False)
add_text(slide, "MiniGUI V1.0.0", 4.88, 4.02, 3.0, 0.35, 13, BLUE, True)
add_text(slide, "2026. 08", 4.88, 4.45, 2.0, 0.3, 12, RGBColor(102, 112, 133))
add_text(slide, "운영 전 본 매뉴얼의 EPO 주의사항을 반드시 확인하십시오.", 4.88, 6.55, 7.3, 0.3, 11, RED, True)

# 2. Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "사용 절차 한눈에 보기", "QUICK START", 2)
add_footer(slide)
steps = [
    (1, "프로그램 실행", "초기 화면과 READY 상태 확인"),
    (2, "설치 정보 입력", "설치 장소와 축전지명 입력 후 저장"),
    (3, "접속 정보 확인", "시스템 IP와 SNMP Port 확인"),
    (4, "시스템 접속", "접속시작 클릭 후 결과 팝업 확인"),
    (5, "운전 상태 확인", "ONLINE, 정상 접속 및 Rack 값 확인"),
    (6, "설정 및 제어", "필요 시 SOC 설정 또는 EPO 수행"),
]
for i, (n, title, body) in enumerate(steps):
    col, row = i % 2, i // 2
    add_step_card(slide, n, title, body, 0.75 + col * 6.15, 1.25 + row * 1.55, 5.7,
                  GREEN if n == 5 else BLUE)
add_rect(slide, 0.75, 6.05, 11.85, 0.62, PALE_RED, RGBColor(254, 205, 202))
add_text(slide, "주의  |  EPO는 전체 모듈 출력을 차단·복구하는 기능입니다. 권한을 가진 운영자만 사용하십시오.",
         1.0, 6.14, 11.35, 0.4, 12, RED, True)

# 3. Initial screen
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "1. 프로그램 실행 및 초기 화면 확인", "START", 3)
add_footer(slide)
add_image_with_callouts(slide, BEFORE, 0.55, 1.08, 6.35, 5.72, [
    (1, (755, 55, 100, 92), BLUE),
    (2, (24, 157, 828, 91), BLUE),
    (3, (24, 367, 828, 79), ORANGE),
])
add_text(slide, "초기 확인 항목", 7.35, 1.25, 4.9, 0.45, 22, NAVY, True)
add_step_card(slide, 1, "READY 확인", "우측 상단 표시가 READY인지 확인합니다.", 7.25, 1.9, 5.25)
add_step_card(slide, 2, "접속 정보 확인", "운영 대상 시스템의 IP와 SNMP Port를 확인합니다.", 7.25, 3.12, 5.25)
add_step_card(slide, 3, "접속 대기 확인", "접속 전 Rack 값이 '-'로 표시되는 것은 정상입니다.", 7.25, 4.34, 5.25, ORANGE)
add_rect(slide, 7.25, 5.76, 5.25, 0.72, PALE_BLUE, RGBColor(180, 205, 255))
add_text(slide, "기본 SNMP Port는 161입니다. 현장 설정이 다르면 지정값을 입력하십시오.",
         7.52, 5.88, 4.72, 0.45, 11, NAVY, True)

# 4. Installation and connection
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "2. 설치 정보 저장 및 시스템 접속", "CONNECT", 4)
add_footer(slide)
add_image_with_callouts(slide, BEFORE, 0.55, 1.08, 6.35, 5.72, [
    (1, (24, 164, 665, 84), BLUE),
    (2, (24, 266, 828, 82), GREEN),
    (3, (699, 178, 136, 58), ORANGE),
])
add_text(slide, "입력 및 접속 순서", 7.35, 1.17, 4.9, 0.45, 22, NAVY, True)
add_step_card(slide, 1, "시스템 접속 정보", "시스템 IP와 SNMP Port가 현장 정보와 일치하는지 확인합니다.", 7.25, 1.82, 5.25)
add_step_card(slide, 2, "설치 정보 저장", "설치 장소와 축전지명을 입력하고 [저장]을 클릭합니다.", 7.25, 3.04, 5.25, GREEN)
add_step_card(slide, 3, "접속 시작", "[접속시작]을 클릭합니다. 결과 팝업 전까지 버튼은 비활성화됩니다.", 7.25, 4.26, 5.25)
add_rect(slide, 7.25, 5.68, 5.25, 0.86, LIGHT_GRAY, MID_GRAY)
add_text(slide, "팝업 확인\n성공: 축전지 시스템 연결 성공  |  실패: IP·Port·통신망 확인",
         7.5, 5.78, 4.78, 0.58, 11, GRAY, True)

# 5. Connected dashboard
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "3. 접속 성공 및 운전 상태 확인", "MONITOR", 5)
add_footer(slide)
add_image_with_callouts(slide, AFTER, 0.55, 1.05, 6.1, 5.95, [
    (1, (598, 45, 86, 74), GREEN),
    (2, (19, 294, 664, 65), GREEN),
    (3, (19, 373, 664, 91), BLUE),
])
add_text(slide, "정상 접속 판단 기준", 7.1, 1.18, 5.0, 0.45, 22, NAVY, True)
add_step_card(slide, 1, "ONLINE", "우측 상단 표시가 녹색 ONLINE으로 변경됩니다.", 7.0, 1.82, 5.55, GREEN)
add_step_card(slide, 2, "정상 접속", "접속 상태에 녹색 점과 '정상 접속'이 표시됩니다.", 7.0, 3.04, 5.55, GREEN)
add_step_card(slide, 3, "Rack 값 수신", "전체용량, 전압, 전류 값과 최종 업데이트 시각을 확인합니다.", 7.0, 4.26, 5.55)
add_rect(slide, 7.0, 5.68, 5.55, 0.82, PALE_RED, RGBColor(254, 205, 202))
add_text(slide, "값이 장시간 갱신되지 않거나 '-'로 유지되면 네트워크 및 장비 상태를 확인하십시오.",
         7.27, 5.82, 5.0, 0.5, 11, RED, True)

# 6. Rack info
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "4. Rack 기본정보 읽기", "MONITOR", 6)
add_footer(slide)
add_image_with_callouts(slide, AFTER, 0.55, 1.08, 6.2, 5.75, [
    (1, (20, 374, 662, 89), BLUE),
    (2, (458, 296, 220, 55), ORANGE),
])
add_text(slide, "1  Rack 기본정보", 7.22, 1.22, 4.8, 0.45, 22, BLUE, True)
add_rect(slide, 7.12, 1.85, 5.5, 2.15, PALE_BLUE, RGBColor(180, 205, 255))
add_text(slide, "전체용량", 7.45, 2.08, 1.35, 0.32, 14, NAVY, True)
add_text(slide, "Rack 정격/설정 용량 (Ah)", 8.85, 2.08, 3.25, 0.32, 13)
add_text(slide, "전압", 7.45, 2.63, 1.35, 0.32, 14, NAVY, True)
add_text(slide, "현재 Rack 전압 (V)", 8.85, 2.63, 3.25, 0.32, 13)
add_text(slide, "전류", 7.45, 3.18, 1.35, 0.32, 14, NAVY, True)
add_text(slide, "현재 Rack 전류 (A)", 8.85, 3.18, 3.25, 0.32, 13)
add_text(slide, "2  업데이트 시각", 7.45, 3.7, 1.35, 0.24, 11, ORANGE, True)
add_text(slide, "수신 데이터의 최신성을 판단합니다.", 8.85, 3.66, 3.25, 0.32, 11)
add_rect(slide, 7.12, 4.35, 5.5, 1.25, LIGHT_GRAY, MID_GRAY)
add_text(slide, "참고", 7.42, 4.55, 0.7, 0.28, 13, NAVY, True)
add_text(slide, "전류 부호는 시스템 운전 방향을 나타냅니다. 현장 운전 기준에 따라 충전/방전 방향을 판단하십시오.",
         8.1, 4.48, 4.15, 0.62, 12, GRAY)

# 7. SOC
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "5. SOC 충전제한 설정", "CONTROL", 7)
add_footer(slide)
add_image_with_callouts(slide, AFTER, 0.55, 1.08, 6.2, 5.75, [
    (1, (88, 485, 150, 58), GREEN),
    (2, (570, 485, 102, 58), BLUE),
])
add_text(slide, "설정 방법", 7.22, 1.24, 4.8, 0.45, 22, NAVY, True)
add_step_card(slide, 1, "현재 설정 확인", "SOC 충전제한 영역의 현재 설정값을 확인합니다.", 7.12, 1.88, 5.5, GREEN)
add_step_card(slide, 2, "설정 변경", "[설정 변경]을 클릭하고 허용된 범위의 값을 입력합니다.", 7.12, 3.10, 5.5)
add_rect(slide, 7.12, 4.32, 5.5, 1.05, PALE_BLUE, RGBColor(180, 205, 255))
add_text(slide, "적용 결과 확인", 7.38, 4.43, 4.98, 0.3, 14, NAVY, True)
add_text(slide, "설정 완료 후 1번 현재 설정값이 변경되었는지 확인합니다.",
         7.38, 4.78, 4.98, 0.38, 11, GRAY)
add_rect(slide, 7.12, 5.72, 5.5, 0.78, PALE_RED, RGBColor(254, 205, 202))
add_text(slide, "운전 정책과 다른 SOC 값은 축전지 운전에 영향을 줄 수 있습니다. 승인된 값만 입력하십시오.",
         7.38, 5.84, 4.98, 0.52, 11, RED, True)

# 8. EPO
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "6. EPO 전체차단 및 전체복구", "EMERGENCY", 8)
add_footer(slide)
add_image_with_callouts(slide, AFTER, 0.55, 1.08, 6.2, 5.75, [
    (1, (402, 566, 126, 48), RED),
    (2, (536, 566, 125, 48), GREEN),
])
add_text(slide, "EPO 기능", 7.22, 1.22, 4.8, 0.45, 22, NAVY, True)
add_rect(slide, 7.12, 1.82, 5.5, 1.02, PALE_RED, RGBColor(254, 205, 202))
add_text(slide, "1  전체차단(모듈:N)", 7.42, 1.96, 4.9, 0.32, 15, RED, True)
add_text(slide, "표시된 N개 모듈에 차단 명령을 순차 전송합니다.", 7.42, 2.33, 4.9, 0.28, 11, GRAY)
add_rect(slide, 7.12, 3.02, 5.5, 1.02, PALE_GREEN, RGBColor(171, 239, 198))
add_text(slide, "2  전체복구(모듈:N)", 7.42, 3.16, 4.9, 0.32, 15, GREEN, True)
add_text(slide, "표시된 N개 인식 모듈에 복구 명령을 순차 전송합니다.", 7.42, 3.53, 4.9, 0.28, 11, GRAY)
add_rect(slide, 7.12, 4.32, 5.5, 1.78, RED, RED)
add_text(slide, "안전 주의", 7.42, 4.52, 4.9, 0.33, 16, WHITE, True)
add_bullet_list(slide, [
    "작업 전 부하와 현장 안전 상태를 확인합니다.",
    "확인 팝업의 대상과 작업 종류를 다시 확인합니다.",
    "긴급 상황 또는 승인된 작업에서만 실행합니다.",
], 7.4, 4.9, 4.9, 1.02, 11, WHITE, 3)

# 9. Disconnect / troubleshooting
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "7. 접속 종료 및 문제 해결", "FINISH", 9)
add_footer(slide)
add_text(slide, "정상 종료", 0.75, 1.25, 5.2, 0.45, 22, NAVY, True)
add_step_card(slide, 1, "제어 작업 완료 확인", "진행 중인 SOC 설정 또는 EPO 작업이 없는지 확인합니다.", 0.7, 1.9, 5.65)
add_step_card(slide, 2, "접속종료 클릭", "시스템 접속 영역의 [접속종료]를 클릭합니다.", 0.7, 3.12, 5.65)
add_step_card(slide, 3, "READY 복귀 확인", "상단 READY와 접속 대기 표시를 확인한 후 프로그램을 종료합니다.", 0.7, 4.34, 5.65)
add_text(slide, "문제 발생 시 확인", 6.95, 1.25, 5.2, 0.45, 22, NAVY, True)
issues = [
    ("접속 실패", "IP, SNMP Port, 통신망 및 대상 장비 전원 확인"),
    ("Rack 값 미표시", "ONLINE/정상 접속 여부와 최종 업데이트 시각 확인"),
    ("버튼 비활성", "접속 상태, 인식 모듈 수 및 진행 중 작업 확인"),
    ("EPO 실패", "실패 모듈 번호 확인 후 현장 담당자에게 전달"),
]
for idx, (title, body) in enumerate(issues):
    y = 1.9 + idx * 1.02
    add_rect(slide, 6.9, y, 5.65, 0.82, LIGHT_GRAY, MID_GRAY)
    add_text(slide, title, 7.18, y + 0.12, 1.25, 0.3, 13, NAVY, True)
    add_text(slide, body, 8.45, y + 0.08, 3.8, 0.5, 11, GRAY)
add_rect(slide, 6.9, 6.05, 5.65, 0.52, PALE_BLUE, RGBColor(180, 205, 255))
add_text(slide, "문의 시 화면 캡처, 발생 시각, IP, 실패 모듈 번호를 함께 전달하십시오.",
         7.15, 6.12, 5.15, 0.3, 10, NAVY, True)


# Document metadata
prs.core_properties.title = "TBC1000B Battery Control Mini 고객 사용자 매뉴얼"
prs.core_properties.subject = "시스템 접속, 상태 확인, SOC 설정 및 EPO 사용 방법"
prs.core_properties.author = "PANTECH C&I Engineering"
prs.core_properties.keywords = "TBC1000B, MiniGUI, 사용자 매뉴얼, EPO, SOC"
prs.save(OUTPUT)
print(OUTPUT)
