import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from photoword.core import Element, Page
from photoword.exporters import FORMATS, export_bytes, save_export
from photoword.insights import Table, detect_tables, recommend


def grid_page():
    image = Image.new("RGB", (600, 800), "white")
    draw = ImageDraw.Draw(image)
    for x in (50, 250, 450):
        draw.line((x, 100, x, 400), fill="black", width=3)
    for y in (100, 250, 400):
        draw.line((50, y, 450, y), fill="black", width=3)
    words = [Element("word", (80,150,60,30),"이름"), Element("word",(280,150,60,30),"수량"),
             Element("word",(80,300,60,30),"사과"), Element("word",(280,300,60,30),"0012")]
    return Page(image, [Element("text",e.box,e.text,99) for e in words], words=words)


class FormatTests(unittest.TestCase):
    def test_table_coordinates_assign_words_to_cells(self):
        tables = detect_tables(grid_page())
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0].cells,[["이름","수량"],["사과","0012"]])

    def test_recommendations_follow_content_and_are_unique(self):
        page = grid_page()
        self.assertEqual(recommend([page])[0].format,"xlsx")
        document = Page(Image.new("RGB",(600,800),"white"),[Element("text",(10,10,400,40),"설명 문장입니다. "*20)])
        self.assertEqual(recommend([document])[0].format,"docx")
        document.elements.extend([Element("rect",(10,100,100,100)),Element("oval",(200,100,100,100))])
        self.assertEqual(recommend([document])[0].format,"pptx")
        document.ocr_available = False
        result = recommend([document])
        self.assertEqual(result[0].format,"pdf")
        self.assertEqual(len({r.format for r in result}),3)

    def test_xlsx_contains_real_cells_and_images_without_formulas(self):
        page = grid_page()
        table = Table((50,100,400,300), [["=1+1","0012"],["한글","두 줄\n값"]])
        workbook = load_workbook(io.BytesIO(export_bytes([page],"xlsx",tables=[[table]])))
        self.assertEqual(workbook["P1_표1"]["A1"].value,"=1+1")
        self.assertEqual(workbook["P1_표1"]["A1"].data_type,"s")
        self.assertEqual(workbook["P1_표1"]["B1"].value,"0012")
        self.assertEqual(len(workbook["P1_원본과텍스트"]._images),1)

    def test_pptx_contains_editable_text_shapes_and_two_slides(self):
        page = grid_page()
        page.elements.extend([Element("rect",(10,500,100,100)),Element("oval",(200,500,100,100)),Element("image",(400,500,100,100))])
        presentation = Presentation(io.BytesIO(export_bytes([page,page],"pptx")))
        self.assertEqual(len(presentation.slides),2)
        self.assertTrue(any(getattr(shape,"text","") == "이름" for shape in presentation.slides[0].shapes))
        self.assertEqual(len(presentation.slides[0].shapes),7)

    def test_csv_roundtrip_and_formula_neutralization(self):
        table = Table((0,0,10,10), [["=1+1","한글, 쉼표"],["두\n줄","0012"]])
        data = export_bytes([grid_page()],"csv",tables=[[table]]).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(data)))
        self.assertEqual(rows[1][3:], ["'=1+1","한글, 쉼표"])
        self.assertEqual(rows[2][3:], ["두\n줄","0012"])
        self.assertEqual(len(rows[0]),len(rows[1]))

    def test_txt_and_pdf_and_no_ocr_fallback(self):
        page = grid_page()
        self.assertIn("사과",export_bytes([page],"txt").decode("utf-8-sig"))
        pdf = export_bytes([page,page],"pdf")
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertIn(b"/Count 2",pdf)
        blank = Page(page.image, [Element("image",(0,0,*page.image.size))],ocr_available=False)
        for format in FORMATS:
            self.assertGreater(len(export_bytes([blank],format)),10)

    def test_failed_export_does_not_damage_existing_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/"existing.docx"
            path.write_bytes(b"original")
            with patch("photoword.exporters.export_bytes",side_effect=RuntimeError("failed")):
                with self.assertRaises(RuntimeError):
                    save_export(path,[grid_page()],"docx")
            self.assertEqual(path.read_bytes(),b"original")

    def test_invalid_extension_and_empty_input(self):
        with self.assertRaises(ValueError):
            save_export("wrong.txt",[grid_page()],"xlsx")
        with self.assertRaises(ValueError):
            export_bytes([],"docx")


if __name__ == "__main__":
    unittest.main()
