"""Single-file Windows launcher and offline packaging verification."""
import os
import sys
import runpy


def self_test(report):
    import json
    import tempfile
    import traceback
    from pathlib import Path
    result = {'ok': False}
    try:
        os.environ['QT_QPA_PLATFORM'] = 'windows' if os.name == 'nt' else 'offscreen'
        import monitor
        from PySide6.QtWidgets import QApplication
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QPixmap, QIcon
        from pysnmp.hlapi import SnmpEngine
        from openpyxl import Workbook, load_workbook
        from upload_transport import Outbox
        import ssl
        app = QApplication([])
        for asset in ('install_battery.png', 'skt_logo.png', 'pantech.png'):
            assert not QPixmap(monitor.resource_path(asset)).isNull(), asset
        assert not QIcon(monitor.resource_path('battery#3.ico')).isNull()
        assert Path(monitor.resource_path('alarm.wav')).stat().st_size > 0
        engine = SnmpEngine()
        engine.getMibBuilder().importSymbols('SNMPv2-MIB', 'sysDescr')
        with tempfile.TemporaryDirectory() as folder:
            dialog = monitor.ProfileDialog(str(Path(folder) / 'profiles'))
            dialog.setAttribute(Qt.WA_DontShowOnScreen, True)
            dialog.show()
            app.processEvents()
            assert dialog.isVisible()
            dialog.grab().save(str(Path(report).with_suffix('.png')))
            dialog.close()
            book = Workbook(); book.active['A1'] = '배포 확인'
            book.save(Path(folder) / 'test.xlsx')
            assert load_workbook(Path(folder) / 'test.xlsx').active['A1'].value == '배포 확인'
            Outbox(Path(folder) / 'queue.sqlite3')
        ssl.create_default_context()
        assert monitor.DEVICE_DEFAULTS == {}, 'Local site defaults must not be bundled'
        result = {'ok': True, 'frozen': bool(getattr(sys, 'frozen', False)),
                  'checks': ['Qt profile dialog', 'icons and images', 'alarm audio resource',
                             'SNMP engine and dynamic MIB', 'Excel read/write', 'SQLite queue', 'TLS context', 'clean defaults']}
    except Exception:
        result['error'] = traceback.format_exc()
    Path(report).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if result['ok'] else 1


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--self-test':
        sys.exit(self_test(os.path.abspath(sys.argv[2])))
    runpy.run_module('monitor', run_name='__main__')
