"""Position editable VML text/shapes and image crops on Word pages."""
import io
from statistics import median

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from .core import Page

V = "urn:schemas-microsoft-com:vml"


def estimated_font_size(box_height: int, scale: float) -> float:
    return min(36, max(7, box_height * scale * 1.15))


def v(tag, **attributes):
    node = etree.Element(f"{{{V}}}{tag}", nsmap={"v": V})
    for key, value in attributes.items():
        node.set(key, str(value))
    return node


def build_docx(pages: list[Page], original: bool = False, plain_text: bool = False, font_size_mode: str = "preserve") -> bytes:
    if not pages:
        raise ValueError("변환할 사진이 없습니다.")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    section.top_margin = section.bottom_margin = Pt(24)
    section.left_margin = section.right_margin = Pt(24)
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    uniform_sizes = []
    for page in pages:
        page_scale = min((595.28-48)/page.image.width, (841.89-60)/page.image.height)
        uniform_sizes.extend(
            estimated_font_size(element.box[3], page_scale)
            for element in page.elements
            if element.kind in ("text", "embedded_text")
        )
    uniform_size = median(uniform_sizes) if uniform_sizes else 10
    serial = 0
    for page_index, page in enumerate(pages):
        scale = min((595.28-48)/page.image.width, (841.89-60)/page.image.height)
        if plain_text and not original:
            text_elements = [element for element in page.elements if element.kind in ("text", "embedded_text")]
            text_elements.sort(key=lambda element: (element.box[1], element.box[0]))
            for element_index, element in enumerate(text_elements):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.page_break_before = page_index > 0 and element_index == 0
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run("".join(c for c in element.text if c in "\t\n\r" or ord(c) >= 32))
                run.font.name = "Malgun Gothic"
                font_size = uniform_size if font_size_mode == "uniform" else estimated_font_size(element.box[3], scale)
                run.font.size = Pt(font_size)
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            if not text_elements:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.page_break_before = page_index > 0
            continue
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = page_index > 0
        p.paragraph_format.line_spacing = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        elements = page.elements
        if original:
            from .core import Element
            elements = [Element("image", (0, 0, *page.image.size))]
        # Background graphics first, editable text in front.
        for e in sorted(elements, key=lambda item: item.kind == "text"):
            if e.kind == "embedded_text":
                continue
            serial += 1
            x, y, w, h = e.box
            height = max(h*scale*1.35, 8) if e.kind == "text" else h*scale
            css = (f"position:absolute;margin-left:{24+x*scale:.2f}pt;"
                   f"margin-top:{24+y*scale:.2f}pt;width:{w*scale:.2f}pt;"
                   f"height:{height:.2f}pt;mso-position-horizontal-relative:page;"
                   f"mso-position-vertical-relative:page;z-index:{serial}")
            shape = v("oval" if e.kind == "oval" else "rect", id=f"PhotoWord{serial}", style=css)
            pict = OxmlElement("w:pict")
            pict.append(shape)
            p.add_run()._r.append(pict)
            if e.kind == "text":
                shape.set("filled", "f")
                shape.set("stroked", "f")
                box = v("textbox", inset="0,0,0,0", style="mso-fit-shape-to-text:f;mso-wrap-style:none")
                content = OxmlElement("w:txbxContent")
                para = OxmlElement("w:p")
                props = OxmlElement("w:pPr")
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:before"), "0")
                spacing.set(qn("w:after"), "0")
                props.append(spacing)
                para.append(props)
                run = OxmlElement("w:r")
                rpr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                for key in ("ascii", "hAnsi", "eastAsia"):
                    fonts.set(qn(f"w:{key}"), "맑은 고딕")
                rpr.append(fonts)
                size_element = OxmlElement("w:sz")
                font_size = uniform_size if font_size_mode == "uniform" else estimated_font_size(h, scale)
                size_value = str(round(font_size * 2))
                size_element.set(qn("w:val"), size_value)
                rpr.append(size_element)
                complex_size = OxmlElement("w:szCs")
                complex_size.set(qn("w:val"), size_value)
                rpr.append(complex_size)
                run.append(rpr)
                text = OxmlElement("w:t")
                text.set(qn("xml:space"), "preserve")
                # XML 1.0 cannot contain OCR control characters.
                text.text = "".join(c for c in e.text if c in "\t\n\r" or ord(c) >= 32)
                run.append(text)
                para.append(run)
                content.append(para)
                box.append(content)
                shape.append(box)
            elif e.kind == "image":
                shape.set("stroked", "f")
                blob = io.BytesIO()
                page.image.crop((x, y, x+w, y+h)).save(blob, format="PNG")
                blob.seek(0)
                rid, _ = doc.part.get_or_add_image(blob)
                image_data = v("imagedata")
                image_data.set(qn("r:id"), rid)
                shape.append(image_data)
            else:
                shape.set("strokecolor", f"#{e.color}")
                shape.set("strokeweight", "1pt")
                shape.set("fillcolor", f"#{e.fill or 'FFFFFF'}")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
