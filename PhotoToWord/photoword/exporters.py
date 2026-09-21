"""Real DOCX, XLSX, PPTX, TXT, PDF and CSV writers."""
import csv
import io
import os
import re
import tempfile
from pathlib import Path

from .core import Element
from .document import build_docx
from .insights import detect_tables, text_elements

FORMATS = {
    "docx": ("Word 문서 (.docx)", "편집 가능한 글자·도형과 이미지. OCR이 없으면 원본 사진을 삽입합니다."),
    "xlsx": ("Excel 통합 문서 (.xlsx)", "구분선이 있는 표는 셀로 추출하고, 원본 이미지와 전체 OCR 텍스트도 별도 시트에 넣습니다."),
    "pptx": ("PowerPoint 프레젠테이션 (.pptx)", "사진마다 슬라이드 하나. 인식된 글자·단순 도형은 편집 가능합니다."),
    "txt": ("텍스트 파일 (.txt)", "OCR 글자만 저장합니다. 사진·도형·배치는 포함되지 않습니다."),
    "pdf": ("PDF 문서 (.pdf)", "사진의 모습을 페이지별로 보존합니다. 글자 검색·수정 기능은 없습니다."),
    "csv": ("CSV 데이터 (.csv)", "표를 행 단위로 저장합니다. 표가 없으면 OCR 줄 목록을 저장하며 사진은 제외합니다."),
}


def clean(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def image_stream(image):
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    return stream


def _xlsx(pages, tables):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (page, group) in enumerate(zip(pages, tables), 1):
        sheet = workbook.create_sheet(f"P{index}_원본과텍스트")
        sheet.cell(1, 1, "원본 이미지 / OCR 텍스트 (수치·수식 자동 변환 없음)")
        picture = SheetImage(image_stream(page.image))
        picture.width, picture.height = 480, 480*page.image.height/page.image.width
        sheet.add_image(picture, "D2")
        sheet.column_dimensions["A"].width = 70
        lines = text_elements(page)
        for row, element in enumerate(lines, 3):
            cell = sheet.cell(row, 1, clean(element.text))
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True)
        if not lines:
            sheet.cell(3, 1, "인식된 텍스트 없음. OCR 설정을 확인하세요.")
        for number, table in enumerate(group, 1):
            sheet = workbook.create_sheet(f"P{index}_표{number}")
            for r, row in enumerate(table.cells, 1):
                for c, text in enumerate(row, 1):
                    cell = sheet.cell(r, c, clean(text))
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if r == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="E7EFF9")
                    sheet.column_dimensions[get_column_letter(c)].width = 22
            sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pptx(pages, original):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    presentation = Presentation()
    landscape = sum(p.image.width > p.image.height for p in pages) > len(pages)/2
    sw, sh = (13.333, 7.5) if landscape else (8.267, 11.693)
    presentation.slide_width, presentation.slide_height = Inches(sw), Inches(sh)
    for page in pages:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        scale = min((sw-0.4)/page.image.width, (sh-0.4)/page.image.height)
        ox, oy = (sw-page.image.width*scale)/2, (sh-page.image.height*scale)/2
        elements = page.elements
        if original or not page.ocr_available or not elements:
            elements = [Element("image", (0, 0, *page.image.size))]
        for e in sorted(elements, key=lambda item: item.kind == "text"):
            if e.kind not in ("image", "text", "rect", "oval"):
                continue
            x, y, w, h = e.box
            left, top, width, height = map(Inches, (ox+x*scale, oy+y*scale, w*scale, h*scale))
            if e.kind == "image":
                slide.shapes.add_picture(image_stream(page.image.crop((x,y,x+w,y+h))), left, top, width, height)
            elif e.kind == "text":
                shape = slide.shapes.add_textbox(left, top, width, round(height*1.4))
                frame = shape.text_frame
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.word_wrap = False
                frame.text = clean(e.text)
                for paragraph in frame.paragraphs:
                    paragraph.font.name = "맑은 고딕"
                    paragraph.font.size = Pt(max(5, h*scale*72*0.85))
                    paragraph.font.color.rgb = RGBColor(0,0,0)
            else:
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if e.kind == "oval" else MSO_SHAPE.RECTANGLE,
                                               left, top, width, height)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(e.fill or "FFFFFF")
                shape.line.color.rgb = RGBColor.from_string(e.color)
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _csv_safe(text):
    text = clean(str(text))
    return "'"+text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def export_bytes(pages, format, *, original=False, plain_text=False, font_size_mode="preserve", tables=None):
    if not pages:
        raise ValueError("먼저 사진을 분석해 주세요.")
    if format not in FORMATS:
        raise ValueError("지원하지 않는 출력 형식입니다.")
    tables = tables if tables is not None else [detect_tables(p) for p in pages]
    if len(tables) != len(pages):
        raise ValueError("페이지와 표 분석 결과가 일치하지 않습니다.")
    if format == "docx":
        return build_docx(pages, original, plain_text=plain_text, font_size_mode=font_size_mode)
    if format == "xlsx":
        return _xlsx(pages, tables)
    if format == "pptx":
        return _pptx(pages, original)
    if format == "pdf":
        stream = io.BytesIO()
        pages[0].image.save(stream, "PDF", save_all=True, append_images=[p.image for p in pages[1:]], resolution=150)
        return stream.getvalue()
    if format == "txt":
        content = []
        for index, page in enumerate(pages, 1):
            lines = [clean(e.text) for e in text_elements(page)]
            content.append(f"--- {index}페이지 ---\n" + ("\n".join(lines) or "[인식된 텍스트 없음. OCR 설정을 확인하세요.]"))
        return "\n\n".join(content).encode("utf-8-sig")
    rows = []
    for index, (page, group) in enumerate(zip(pages, tables), 1):
        if group:
            for number, table in enumerate(group, 1):
                rows.extend([index, number, r, *map(_csv_safe, row)] for r, row in enumerate(table.cells, 1))
        else:
            lines = [e.text for e in text_elements(page)] or ["인식된 텍스트 없음"]
            rows.extend([index, 0, r, _csv_safe(line)] for r, line in enumerate(lines, 1))
    columns = max(len(row) for row in rows)
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["페이지", "표번호(0=본문)", "행", *[f"열{i}" for i in range(1, columns-2)]])
    writer.writerows(row+[""]*(columns-len(row)) for row in rows)
    return stream.getvalue().encode("utf-8-sig")


def save_export(path, pages, format, *, original=False, plain_text=False, font_size_mode="preserve", tables=None):
    path = Path(path)
    if path.suffix.lower() != "."+format:
        raise ValueError(f"파일 확장자는 .{format}이어야 합니다.")
    blob = export_bytes(pages, format, original=original, plain_text=plain_text, font_size_mode=font_size_mode, tables=tables)
    # Only replace the destination after a complete successful conversion.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".photoword-", suffix=".tmp", delete=False) as stream:
            temporary = stream.name
            stream.write(blob)
        os.replace(temporary, path)
    finally:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
