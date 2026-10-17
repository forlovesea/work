"""Local, extractive meeting summaries and complete transcript exports."""
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import re
import threading
from xml.sax.saxutils import escape

KINDS = {"speech": "발언", "note": "추가 메모", "decision": "확정 결정", "action": "실행 항목"}
STOP = set("회의 오늘 이번 저희 우리 제가 저는 대한 대해서 관련 그리고 하지만 그래서 그러면 일단 지금 말씀 내용 부분 경우 필요 확인 검토 진행 예정 생각 문제 있습니다 없습니다 합니다 입니다 같습니다 하는 하고 해서 하면 것을 수 것 좀 더 안 네 예 그 이 저 및 또는 아직 우선 계속 다시 모두 정말 매우 어떻게 무엇 어떤 것으로 하겠습니다 대해서는 여기 거기 대한 한다 된다 했다 있다 없다 필요합니다 확인하겠습니다 검토하겠습니다 좋겠습니다 같습니다 대해서 논의 이야기".split())
STOP.update("현재 다음 먼저 함께 전에 이후 이전 결과 상태 해야 확인해야 확인하고 확인한 변경하기 결정하지 않았습니다 됩니다 대한 위한 위한것 위해 때까지 하기로 하도록 통해 따라서 또한 하나 하는데 대해서도".split())
FONT_DIR = Path(__file__).resolve().parent / "fonts"
_font_lock = threading.Lock()


def clean(text):
    # XML 1.0 cannot encode control characters; preserve tabs and line breaks.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", str(text))


def stamp(value):
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(
            timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return clean(value)


def terms(text):
    result = set()
    for token in re.findall(r"[가-힣]{2,}|[A-Za-z][A-Za-z0-9_-]+", text):
        token = re.sub(r"(에서는|으로는|에게는|에서도|에서|으로|에게|까지|부터|에는|은|는|이|가|을|를|의|와|과|도)$", "", token)
        if len(token) >= 2 and token not in STOP and not token.endswith(("니다", "세요", "어요", "아요")):
            result.add(token)
    return result


def summary_blocks(entries):
    blocks = [("h1", "회의록 요약"), ("note", "전체 저장 기록에서 주요 키워드와 대표 발언을 발췌한 요약입니다. 자동 전사의 오인식과 생략된 맥락은 아래 전체 대화에서 확인하세요.")]
    sentences = []
    seen = set()
    for e in entries:
        if e["kind"] != "speech":
            continue
        # Keep each utterance together so follow-up sentences retain their topic.
        sentence = clean(e["text"]).strip()
        if sentence and sentence not in seen:
            seen.add(sentence)
            sentences.append((sentence, terms(sentence), e))
    frequency = Counter(t for _, ts, _ in sentences for t in ts)
    first_seen = {}
    for i, (text, tokens, _) in enumerate(sentences):
        for token in tokens:
            first_seen.setdefault(token, (i, text.find(token)))
    remaining = set(range(len(sentences)))
    blocks.append(("h2", "주요 주제와 내용"))
    if not sentences:
        blocks.append(("body", "저장된 음성 발언이 없습니다. 아래 방장 기록을 확인하세요."))
    for _ in range(6):
        counts = Counter(t for i in sorted(remaining) for t in sentences[i][1])
        if not counts:
            break
        topic = sorted(counts, key=lambda t: (-counts[t], -frequency[t], first_seen[t], t))[0]
        matches = [i for i in sorted(remaining) if topic in sentences[i][1]]
        # Prefer sentences covering recurring terms, retain source order for reading.
        ranked = sorted(matches, key=lambda i: (-sum(frequency[t] for t in sentences[i][1]) / max(1, len(sentences[i][1])), i))[:3]
        related = Counter(t for i in matches for t in sentences[i][1] if t != topic)
        companion = sorted(related, key=lambda t: (-related[t], first_seen[t], t))[:1]
        blocks.append(("h3", " · ".join([topic, *companion])))
        for i in sorted(ranked):
            text, _, e = sentences[i]
            excerpt = text if len(text) <= 400 else text[:400] + "…"
            blocks.append(("body", f"{excerpt}\n출처: {e['speaker']} · {stamp(e['at'])}"))
        remaining.difference_update(matches)
    if sentences and not frequency:
        for text, _, e in sentences[:3]:
            blocks.append(("body", f"{text[:400]} — {e['speaker']}"))
    for kind, heading in [("decision", "확정 결정"), ("action", "실행 항목"), ("note", "추가 메모")]:
        blocks.append(("h2", heading))
        selected = [e for e in entries if e["kind"] == kind]
        blocks.extend(("body", f"{e['text']}\n기록: {e['speaker']} · {stamp(e['at'])}") for e in selected)
        if not selected:
            blocks.append(("body", "방장이 확정하여 기록한 항목 없음" if kind != "note" else "추가 메모 없음"))
    return blocks


def document_blocks(title, created_at, entries, include_summary=False):
    entries = sorted(entries, key=lambda e: (e["at"], e.get("id", 0)))
    speakers = list(dict.fromkeys(e["speaker"] for e in entries))
    blocks = [("title", f"{title} 회의록"), ("note", f"회의 시작: {stamp(created_at)} KST\n저장 기록: {len(entries)}건 · 발언 및 기록자: {', '.join(speakers) or '없음'}"),
              ("note", "자동 기록이 켜진 동안 서버에 저장된 전체 발언과 수동 기록을 수록합니다. 음성 인식 결과는 실제 발언과 다를 수 있습니다.")]
    if include_summary:
        blocks.extend(summary_blocks(entries))
    blocks.append(("h1", "전체 대화 및 기록"))
    if not entries:
        blocks.append(("body", "저장된 대화 및 기록이 없습니다."))
    for e in entries:
        blocks.append(("speaker", f"{stamp(e['at'])} KST | {e['speaker']} | {KINDS.get(e['kind'], e['kind'])}"))
        blocks.append(("body", e["text"]))
    return [(kind, clean(text)) for kind, text in blocks]


def export_minutes(title, created_at, entries, file_format="txt", include_summary=False):
    blocks = document_blocks(title, created_at, entries, include_summary)
    if file_format == "txt":
        return ("\n\n".join(text for kind, text in blocks if kind != "break") + "\n").encode("utf-8-sig"), "text/plain; charset=utf-8"
    if file_format == "docx":
        return word_bytes(blocks), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if file_format == "pdf":
        return pdf_bytes(blocks), "application/pdf"
    raise ValueError("Unsupported export format")


def word_bytes(blocks):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name, size in [("Normal", 10.5), ("Title", 22), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Malgun Gothic", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = Pt(size * 1.5)
    doc.styles["Title"].font.bold = True
    styles = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}
    for kind, text in blocks:
        if kind == "break":
            doc.add_page_break()
            continue
        p = doc.add_paragraph(text, styles.get(kind, "Normal"))
        if kind == "speaker":
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(9)
        elif kind == "note":
            p.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("페이지 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_bytes(blocks):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
    with _font_lock:
        if "MoaKorean" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("MoaKorean", str(FONT_DIR / "NanumGothic-Regular.ttf")))
            pdfmetrics.registerFont(TTFont("MoaKoreanBold", str(FONT_DIR / "NanumGothic-Bold.ttf")))
    sizes = {"title": 22, "h1": 17, "h2": 13, "h3": 11, "speaker": 9, "note": 9, "body": 10.5}
    styles = {kind: ParagraphStyle(kind, fontName="MoaKoreanBold" if kind in ("title", "h1", "h2", "h3", "speaker") else "MoaKorean",
                                  fontSize=size, leading=size * 1.55, spaceAfter=8, spaceBefore=8 if kind.startswith("h") else 0,
                                  wordWrap="CJK", keepWithNext=kind in ("h1", "h2", "h3", "speaker")) for kind, size in sizes.items()}
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=62, rightMargin=62, topMargin=55, bottomMargin=55,
                            title=blocks[0][1], author="Moa")
    flow = [PageBreak() if kind == "break" else Paragraph(escape(text).replace("\n", "<br/>"), styles[kind]) for kind, text in blocks]
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("MoaKorean", 8)
        canvas.drawRightString(A4[0] - 62, 30, f"페이지 {doc.page}")
        canvas.restoreState()
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
