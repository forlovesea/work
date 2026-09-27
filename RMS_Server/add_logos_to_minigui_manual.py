from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "TBC1000B_MiniGUI_고객사용자매뉴얼(260810).pptx"
SKT_LOGO = ROOT / "skt_logo-original.png"
PANTECH_LOGO = ROOT / "pantech-original.png"
FONT = "맑은 고딕"
WHITE = RGBColor(255, 255, 255)
MID_GRAY = RGBColor(208, 213, 221)
LABEL_GRAY = RGBColor(152, 162, 179)


def add_rect(slide, name, x, y, w, h, fill, line, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_label(slide, name, value, x, y, w, h):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = frame.paragraphs[0].add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = LABEL_GRAY
    return shape


def add_picture_contain(slide, name, path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    px, py = x + (w - dw) / 2, y + (h - dh) / 2
    picture = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(dw), Inches(dh))
    picture.name = name


def add_line(slide, name, x1, y1, x2, y2, color, width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.name = name
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


prs = Presentation(str(PPTX))
cover = prs.slides[0]

managed_names = {
    "Logo partner panel", "Logo divider", "Customer label", "Engineering label",
    "SKT customer logo", "Pantech engineering logo",
    "Logo line top", "Logo line vertical", "Logo line bottom",
    "Logo point top", "Logo point bottom",
    "SKT logo card", "SKT accent red", "SKT accent orange",
    "Pantech logo card", "Pantech accent",
    "SKT logo underline", "Pantech logo overline",
}
for shape in list(cover.shapes):
    if shape.name in managed_names:
        cover.shapes._spTree.remove(shape._element)

# 고객사와 제작사를 서로 독립된 브랜드 카드로 표현한다.
add_rect(cover, "SKT logo card", 9.76, 0.085, 2.74, 1.28, WHITE, MID_GRAY)
add_rect(cover, "SKT accent red", 0.00, 0.00, 10.65, 0.085, RGBColor(234, 0, 44), RGBColor(234, 0, 44), False)
add_rect(cover, "SKT accent orange", 10.65, 0.00, 2.683, 0.085, RGBColor(255, 112, 0), RGBColor(255, 112, 0), False)
add_rect(cover, "SKT logo underline", 9.58, 0.00, 3.10, 0.22, RGBColor(234, 0, 44), RGBColor(234, 0, 44), True)
add_label(cover, "Customer label", "CUSTOMER", 10.02, 0.23, 1.4, 0.2)
add_picture_contain(cover, "SKT customer logo", SKT_LOGO, 10.00, 0.47, 2.24, 0.72)

add_rect(cover, "Pantech logo card", 9.18, 6.135, 3.32, 1.28, WHITE, MID_GRAY)
add_rect(cover, "Pantech accent", 0.00, 7.415, 13.333, 0.085, RGBColor(0, 143, 174), RGBColor(0, 143, 174), False)
add_rect(cover, "Pantech logo overline", 9.00, 7.28, 3.68, 0.22, RGBColor(0, 143, 174), RGBColor(0, 143, 174), True)
add_label(cover, "Engineering label", "SYSTEM ENGINEERING", 9.46, 6.27, 2.0, 0.2)
add_picture_contain(cover, "Pantech engineering logo", PANTECH_LOGO, 9.43, 6.50, 2.78, 0.65)

# 하단 브랜드 카드와 겹치지 않도록 기존 안전 문구를 왼쪽 여백에 정돈한다.
for shape in cover.shapes:
    if hasattr(shape, "text") and "EPO 주의사항" in shape.text:
        shape.left = Inches(4.88)
        shape.top = Inches(5.66)
        shape.width = Inches(4.05)
        shape.height = Inches(0.55)

prs.save(str(PPTX))
print(PPTX)
